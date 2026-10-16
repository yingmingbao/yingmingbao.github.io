# -*- coding: utf-8 -*-
"""从 Markdown 源生成 HTML / PDF / DOCX 三份成品。"""
import os, re
import markdown

D = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(D, "Web容器与数据库技术史.md")
FONTS = os.path.join(D, "fonts")

# ---- 保留中文的 slugify，供 toc 扩展生成稳定的锚点 ----
def slugify(text, sep="-"):
    text = text.strip()
    s = re.sub(r"[\s]+", sep, text)
    s = re.sub(r"[^\w\u4e00-\u9fff\-]", "", s)  # 保留字母/数字/下划线/中文/连字符
    return s.strip(sep)

md = markdown.Markdown(
    extensions=["tables", "fenced_code", "toc", "sane_lists"],
    extension_configs={"toc": {"slugify": slugify, "toc_depth": 3}},
)
with open(SRC, "r", encoding="utf-8") as f:
    SRC_TEXT = f.read()
HTML_BODY = md.convert(SRC_TEXT)

TITLE = "中间件 Web 容器与数据库软件技术发展史"

# ---- 预览用 HTML（依赖系统 CJK 字体，浏览器可正常显示）----
HTML_CSS = """
* { box-sizing: border-box; }
body { font-family: "Microsoft YaHei","SimSun","Songti SC",sans-serif;
       font-size: 15px; line-height: 1.75; color:#1f2430; max-width: 1080px;
       margin: 0 auto; padding: 32px 28px 64px; }
h1 { font-size: 28px; border-bottom: 3px solid #2b5797; padding-bottom: 10px; color:#1b3a6b; }
h2 { font-size: 22px; margin-top: 38px; color:#2b5797; border-left: 6px solid #2b5797; padding-left: 10px; }
h3 { font-size: 18px; margin-top: 26px; color:#3358a0; }
h4 { font-size: 16px; color:#444; }
a  { color:#2b5797; text-decoration: none; }
code { background:#f2f4f8; padding:1px 5px; border-radius:4px; font-family: Consolas,Menlo,monospace; font-size: 13px; }
pre code { display:block; padding:12px; }
blockquote { background:#f6f8fc; border-left:4px solid #9bb4e0; margin: 14px 0; padding: 8px 14px; color:#334; }
table { border-collapse: collapse; width: 100%; margin: 16px 0; font-size: 13.5px; }
th,td { border: 1px solid #c4ccd8; padding: 7px 9px; vertical-align: top; text-align: left; }
th { background:#2b5797; color:#fff; }
tbody tr:nth-child(even) { background:#eef2f8; }
.toc { background:#f6f8fc; border:1px solid #d4dcea; border-radius:8px; padding:14px 20px; margin: 18px 0; }
.toc ul { margin: 4px 0; padding-left: 22px; }
hr { border:none; border-top:1px solid #dde; margin: 26px 0; }
"""

HTML_DOC = f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{TITLE}</title><style>{HTML_CSS}</style></head>
<body>{HTML_BODY}</body></html>"""

with open(os.path.join(D, "Web容器与数据库技术史.html"), "w", encoding="utf-8") as f:
    f.write(HTML_DOC)
print("HTML written")

# ---- PDF 用 HTML（嵌入本地 CJK 字体，保证 xhtml2pdf 中文渲染）----
PDF_CSS = f"""
@page {{ size: A4; margin: 16mm 15mm; }}
@font-face {{ font-family: STSong; src: url("fonts/STSong.TTF"); }}
@font-face {{ font-family: SimHei; src: url("fonts/simhei.ttf"); }}
* {{ box-sizing: border-box; }}
body {{ font-family: STSong; font-size: 10.5pt; line-height: 1.6; color:#1f2430; }}
h1 {{ font-family: SimHei; font-size: 21pt; color:#1b3a6b; }}
h2 {{ font-family: SimHei; font-size: 15pt; color:#2b5797; margin-top: 18pt; }}
h3 {{ font-family: SimHei; font-size: 12.5pt; color:#3358a0; }}
h4 {{ font-family: SimHei; font-size: 11pt; color:#444; }}
a {{ color:#2b5797; }}
code {{ font-family: Consolas,Menlo,monospace; font-size: 9pt; }}
blockquote {{ border-left: 3px solid #9bb4e0; padding: 4px 10px; color:#334; }}
table, td, th {{ border: 1px solid #c4ccd8; }}
table {{ width: 100%; border-collapse: collapse; margin: 10pt 0; font-size: 8.6pt; }}
th, td {{ padding: 4px 6px; vertical-align: top; text-align: left; word-wrap: break-word; }}
th {{ background:#2b5797; color:#fff; }}
.toc {{ border:1px solid #d4dcea; padding:8pt 14pt; margin: 10pt 0; }}
.toc ul {{ padding-left: 18pt; }}
"""

PDF_DOC = f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8"><style>{PDF_CSS}</style></head>
<body>{HTML_BODY}</body></html>"""

os.chdir(D)  # xhtml2pdf 以 cwd 作为资源读取基准目录，确保相对字体路径可读
from xhtml2pdf import pisa
with open(os.path.join(D, "Web容器与数据库技术史.pdf"), "wb") as out:
    res = pisa.CreatePDF(PDF_DOC, out, path=D)
print("PDF written, errors:", res.err)

# ---- DOCX（python-docx 由 HTML 转换，结构保留）----
import re
from html.parser import HTMLParser
from docx import Document
from docx.oxml.ns import qn

class DocxBuilder(HTMLParser):
    def __init__(self, doc):
        super().__init__(convert_charrefs=True)
        self.doc = doc
        self.runs = []
        self.fmt = {"bold": False, "italic": False, "code": False}
        self.fmt_stack = []
        self.block = None
        self.hlevel = 1
        self.list_style = "List Bullet"
        self.table_rows = None
        self.cur_row = None

    def push_fmt(self):
        self.fmt_stack.append(dict(self.fmt))

    def pop_fmt(self):
        if self.fmt_stack:
            self.fmt = self.fmt_stack.pop()

    def handle_starttag(self, tag, attrs):
        if tag in ("h1", "h2", "h3", "h4"):
            self.block = "h"; self.hlevel = int(tag[1]); self.runs = []
        elif tag == "p":
            self.block = "p"; self.runs = []
        elif tag == "li":
            self.block = "li"; self.runs = []
        elif tag == "blockquote":
            self.block = "quote"; self.runs = []
        elif tag in ("ul", "ol"):
            self.list_style = "List Number" if tag == "ol" else "List Bullet"
        elif tag == "table":
            self.table_rows = []; self.cur_row = None
        elif tag == "tr":
            self.cur_row = []
        elif tag in ("td", "th"):
            self.block = "cell"; self.runs = []
        elif tag in ("strong", "b"):
            self.push_fmt(); self.fmt["bold"] = True
        elif tag == "em":
            self.push_fmt(); self.fmt["italic"] = True
        elif tag == "code":
            self.push_fmt(); self.fmt["code"] = True

    def handle_endtag(self, tag):
        if tag in ("h1", "h2", "h3", "h4"):
            self.doc.add_heading("".join(t for t, _ in self.runs), level=self.hlevel)
            self.runs = []
        elif tag == "p":
            self._flush_para(None)
        elif tag == "li":
            self._flush_para(self.list_style)
        elif tag == "blockquote":
            self._flush_para("Quote")
        elif tag in ("td", "th"):
            if self.cur_row is not None:
                self.cur_row.append(list(self.runs))
            self.runs = []
        elif tag == "tr":
            if self.table_rows is not None and self.cur_row is not None:
                self.table_rows.append(self.cur_row)
            self.cur_row = None
        elif tag == "table":
            self._flush_table()
        elif tag in ("strong", "b", "em", "code"):
            self.pop_fmt()
        elif tag in ("ul", "ol"):
            self.list_style = "List Bullet"

    def handle_data(self, data):
        if self.block in ("h", "p", "li", "quote", "cell"):
            self.runs.append((data, dict(self.fmt)))

    def _apply(self, run):
        run.bold = self._last_fmt.get("bold", False)
        run.italic = self._last_fmt.get("italic", False)
        if self._last_fmt.get("code", False):
            run.font.name = "Consolas"

    def _flush_para(self, style):
        p = self.doc.add_paragraph(style=style) if style else self.doc.add_paragraph()
        for text, f in self.runs:
            if text:
                self._last_fmt = f
                r = p.add_run(text)
                self._apply(r)
        self.runs = []

    def _flush_table(self):
        if not self.table_rows:
            return
        ncols = max((len(r) for r in self.table_rows), default=1)
        table = self.doc.add_table(rows=len(self.table_rows), cols=ncols)
        table.style = "Table Grid"
        for ri, row in enumerate(self.table_rows):
            for ci in range(ncols):
                cell_runs = row[ci] if ci < len(row) else []
                cell = table.cell(ri, ci)
                cell.text = ""
                para = cell.paragraphs[0]
                for text, f in cell_runs:
                    if text:
                        self._last_fmt = f
                        r = para.add_run(text)
                        self._apply(r)
        self.table_rows = None
        self.cur_row = None

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "SimSun"
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
normal.font.size = doc.styles["Normal"].font.size  # keep default size

# 去除 TOC 块，避免 DOCX 中重复目录
body_for_docx = re.sub(r'<div class="toc">.*?</div>', "", HTML_BODY, flags=re.S)
builder = DocxBuilder(doc)
builder.feed(body_for_docx)
docx_path = os.path.join(D, "Web容器与数据库技术史.docx")
doc.save(docx_path)
print("DOCX written")

